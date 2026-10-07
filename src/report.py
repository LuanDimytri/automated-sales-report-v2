"""Cálculo de indicadores e geração de um relatório gerencial de vendas em Excel."""

from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# Paleta simples e consistente para o relatório.
NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "00A6A6"
LIGHT_GRAY = "F3F6F9"
DARK = "243447"
WHITE = "FFFFFF"
GREEN = "70AD47"
ORANGE = "ED7D31"

THIN_BORDER = Border(
    left=Side(style="thin", color="D9E1F2"),
    right=Side(style="thin", color="D9E1F2"),
    top=Side(style="thin", color="D9E1F2"),
    bottom=Side(style="thin", color="D9E1F2"),
)

CURRENCY_FORMAT = 'R$ #,##0.00'
INTEGER_FORMAT = '#,##0'


def calculate_kpis(df: pd.DataFrame) -> dict:
    """Calcula os indicadores principais para o conjunto de vendas."""
    faturamento_total = float(df["total_value"].sum())
    quantidade_total = int(df["quantity"].sum())
    total_vendas = int(len(df))
    ticket_medio = faturamento_total / total_vendas if total_vendas else 0.0
    return {
        "faturamento_total": faturamento_total,
        "quantidade_total": quantidade_total,
        "total_vendas": total_vendas,
        "ticket_medio": ticket_medio,
    }


def _aggregate_sales(
    df: pd.DataFrame, group_column: str, include_quantity: bool = True
) -> pd.DataFrame:
    aggregations = {"total_value": "sum"}
    if include_quantity:
        aggregations["quantity"] = "sum"
    grouped = df.groupby(group_column, as_index=False).agg(aggregations)
    grouped = grouped.rename(columns={"total_value": "faturamento_total"})
    if include_quantity:
        grouped = grouped.rename(columns={"quantity": "quantidade_total"})
    return grouped


def get_sales_by_category(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento, unidades e quantidade de vendas por categoria."""
    result = _aggregate_sales(df, "category")
    counts = df.groupby("category").size().rename("total_vendas")
    return result.join(counts, on="category").sort_values("faturamento_total", ascending=False)


def get_sales_by_seller(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento, unidades e quantidade de vendas por vendedor."""
    result = _aggregate_sales(df, "seller")
    counts = df.groupby("seller").size().rename("total_vendas")
    return result.join(counts, on="seller").sort_values("faturamento_total", ascending=False)


def get_sales_by_region(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento e unidades vendidas por região."""
    return _aggregate_sales(df, "region").sort_values("faturamento_total", ascending=False)


def get_sales_by_payment(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento e quantidade de vendas por forma de pagamento."""
    result = _aggregate_sales(df, "payment_method", include_quantity=False)
    counts = df.groupby("payment_method").size().rename("total_vendas")
    return result.join(counts, on="payment_method").sort_values("faturamento_total", ascending=False)


def get_monthly_trend(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento, unidades e vendas por mês (AAAA-MM)."""
    result = _aggregate_sales(df, "year_month").sort_values("year_month")
    counts = df.groupby("year_month").size().rename("total_vendas")
    return result.join(counts, on="year_month")


def get_sales_by_product(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento, unidades e vendas por produto."""
    result = _aggregate_sales(df, "product")
    counts = df.groupby("product").size().rename("total_vendas")
    return result.join(counts, on="product").sort_values("faturamento_total", ascending=False)


def _format_brl(value: float) -> str:
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _format_int(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def _write_table(writer: pd.ExcelWriter, df: pd.DataFrame, sheet_name: str) -> None:
    """Escreve uma tabela auxiliar com formatação consistente."""
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.book[sheet_name]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(color=WHITE, bold=True)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center")

    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header = str(column_cells[0].value)
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 14), 34)

        if header in {"faturamento_total", "ticket_medio", "unit_price", "total_value"}:
            for cell in column_cells[1:]:
                cell.number_format = CURRENCY_FORMAT
        elif header in {"quantidade_total", "total_vendas", "quantity", "sale_id"}:
            for cell in column_cells[1:]:
                cell.number_format = INTEGER_FORMAT

    # Listras leves nas linhas para facilitar leitura.
    for row in ws.iter_rows(min_row=2):
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)


def _style_dashboard_base(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    ws.print_area = "B2:M44"
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.column_dimensions["A"].width = 3
    for col in range(2, 14):
        ws.column_dimensions[get_column_letter(col)].width = 14

    for row in range(1, 46):
        ws.row_dimensions[row].height = 21


def _merge_title(ws, cell_range: str, value: str, size: int = 20) -> None:
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = value
    cell.font = Font(size=size, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _create_kpi_card(ws, cell_range: str, label: str, value: str, fill_color: str = BLUE) -> None:
    ws.merge_cells(cell_range)
    start = cell_range.split(":")[0]
    cell = ws[start]
    cell.value = f"{label}\n{value}"
    cell.font = Font(size=14, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def _add_chart_title(ws, cell: str, title: str) -> None:
    ws[cell] = title
    ws[cell].font = Font(size=12, bold=True, color=DARK)


def _create_insights(df: pd.DataFrame, kpis: dict) -> list[str]:
    category = get_sales_by_category(df)
    product = get_sales_by_product(df)
    monthly = get_monthly_trend(df)

    insights: list[str] = []
    if not category.empty:
        top_category = category.iloc[0]
        share = top_category["faturamento_total"] / kpis["faturamento_total"] if kpis["faturamento_total"] else 0
        insights.append(
            f"{top_category['category']} lidera o faturamento, com {_format_brl(float(top_category['faturamento_total']))} ({share:.1%} do total)."
        )
    if not product.empty:
        top_product = product.iloc[0]
        insights.append(
            f"{top_product['product']} é o produto de maior faturamento, gerando {_format_brl(float(top_product['faturamento_total']))}."
        )
    if not monthly.empty:
        best_month = monthly.loc[monthly["faturamento_total"].idxmax()]
        insights.append(
            f"O maior faturamento mensal ocorreu em {best_month['year_month']}: {_format_brl(float(best_month['faturamento_total']))}."
        )
    insights.append(f"O ticket médio do período foi de {_format_brl(kpis['ticket_medio'])} por venda.")
    return insights


def _build_dashboard(writer: pd.ExcelWriter, df: pd.DataFrame) -> None:
    """Cria a aba executiva do relatório."""
    ws = writer.book.create_sheet("Dashboard", 0)
    _style_dashboard_base(ws)

    kpis = calculate_kpis(df)
    _merge_title(ws, "B2:M3", "RELATÓRIO GERENCIAL DE VENDAS", 20)

    start_date = pd.to_datetime(df["sale_date"]).min().strftime("%d/%m/%Y")
    end_date = pd.to_datetime(df["sale_date"]).max().strftime("%d/%m/%Y")
    ws.merge_cells("B4:M4")
    ws["B4"] = f"Período analisado: {start_date} a {end_date}  |  Dados processados automaticamente"
    ws["B4"].font = Font(italic=True, color="666666")

    _create_kpi_card(ws, "B6:D8", "FATURAMENTO TOTAL", _format_brl(kpis["faturamento_total"]), BLUE)
    _create_kpi_card(ws, "E6:G8", "TOTAL DE VENDAS", _format_int(kpis["total_vendas"]), TEAL)
    _create_kpi_card(ws, "H6:J8", "ITENS VENDIDOS", _format_int(kpis["quantidade_total"]), GREEN)
    _create_kpi_card(ws, "K6:M8", "TICKET MÉDIO", _format_brl(kpis["ticket_medio"]), ORANGE)

    monthly = get_monthly_trend(df)
    category = get_sales_by_category(df)
    products = get_sales_by_product(df).head(10).sort_values("faturamento_total")
    region = get_sales_by_region(df)

    # Dados auxiliares ocultos/afastados para alimentar os gráficos.
    monthly_chart_df = monthly[["year_month", "faturamento_total"]]
    monthly_chart_df.to_excel(writer, sheet_name="_dados_dashboard", index=False, startrow=0)
    category[["category", "faturamento_total"]].to_excel(writer, sheet_name="_dados_dashboard", index=False, startrow=16)
    products[["product", "faturamento_total"]].to_excel(writer, sheet_name="_dados_dashboard", index=False, startrow=28)
    region[["region", "faturamento_total"]].to_excel(writer, sheet_name="_dados_dashboard", index=False, startrow=42)
    hidden = writer.book["_dados_dashboard"]
    hidden.sheet_state = "hidden"

    _add_chart_title(ws, "B10", "Evolução do faturamento")
    line = LineChart()
    line.title = None
    line.y_axis.title = "Faturamento (R$)"
    line.x_axis.title = "Mês"
    line.height = 7.2
    line.width = 14.5
    line.style = 13
    data = Reference(hidden, min_col=2, min_row=1, max_row=1 + len(monthly_chart_df))
    cats = Reference(hidden, min_col=1, min_row=2, max_row=1 + len(monthly_chart_df))
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    line.legend = None
    ws.add_chart(line, "B11")

    _add_chart_title(ws, "H10", "Faturamento por categoria")
    bar = BarChart()
    bar.type = "bar"
    bar.style = 10
    bar.height = 7.2
    bar.width = 11.5
    bar.legend = None
    bar.y_axis.title = "Categoria"
    bar.x_axis.title = "Faturamento (R$)"
    cat_data = Reference(hidden, min_col=2, min_row=17, max_row=17 + len(category))
    cat_labels = Reference(hidden, min_col=1, min_row=18, max_row=17 + len(category))
    bar.add_data(cat_data, titles_from_data=True)
    bar.set_categories(cat_labels)
    ws.add_chart(bar, "H11")

    _add_chart_title(ws, "B25", "Top 10 produtos por faturamento")
    product_bar = BarChart()
    product_bar.type = "bar"
    product_bar.style = 12
    product_bar.height = 7.2
    product_bar.width = 14.5
    product_bar.legend = None
    prod_data = Reference(hidden, min_col=2, min_row=29, max_row=29 + len(products))
    prod_labels = Reference(hidden, min_col=1, min_row=30, max_row=29 + len(products))
    product_bar.add_data(prod_data, titles_from_data=True)
    product_bar.set_categories(prod_labels)
    ws.add_chart(product_bar, "B26")

    _add_chart_title(ws, "H25", "Participação do faturamento por região")
    doughnut = DoughnutChart()
    doughnut.height = 7.2
    doughnut.width = 11.5
    region_data = Reference(hidden, min_col=2, min_row=43, max_row=43 + len(region))
    region_labels = Reference(hidden, min_col=1, min_row=44, max_row=43 + len(region))
    doughnut.add_data(region_data, titles_from_data=True)
    doughnut.set_categories(region_labels)
    doughnut.holeSize = 55
    doughnut.dataLabels = DataLabelList()
    doughnut.dataLabels.showPercent = True
    doughnut.dataLabels.showVal = False
    doughnut.dataLabels.showCatName = False
    doughnut.dataLabels.showSerName = False
    ws.add_chart(doughnut, "H26")

    # Insights text.
    ws.merge_cells("B40:M40")
    ws["B40"] = "PRINCIPAIS INSIGHTS"
    ws["B40"].font = Font(size=12, bold=True, color=WHITE)
    ws["B40"].fill = PatternFill("solid", fgColor=NAVY)
    ws["B40"].alignment = Alignment(horizontal="left")

    for index, insight in enumerate(_create_insights(df, kpis), start=41):
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=13)
        ws.cell(index, 2).value = f"• {insight}"
        ws.cell(index, 2).font = Font(size=10, color=DARK)
        ws.cell(index, 2).alignment = Alignment(wrap_text=True, vertical="center")


def _format_workbook(writer: pd.ExcelWriter) -> None:
    """Aplica formatação adicional às planilhas auxiliares."""
    for worksheet in writer.book.worksheets:
        if worksheet.title in {"Dashboard", "_dados_dashboard"}:
            continue
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions


def generate_excel_report(df: pd.DataFrame, output_path: str | Path) -> None:
    """Gera o relatório gerencial completo em Excel."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Dados e análises detalhadas continuam disponíveis para auditoria.
        _write_table(writer, get_sales_by_category(df), "Categorias")
        _write_table(writer, get_sales_by_product(df), "Produtos")
        _write_table(writer, get_sales_by_seller(df), "Vendedores")
        _write_table(writer, get_sales_by_region(df), "Regiões")
        _write_table(writer, get_sales_by_payment(df), "Pagamentos")
        _write_table(writer, get_monthly_trend(df), "Evolução Mensal")
        _write_table(writer, df, "Dados")

        _build_dashboard(writer, df)
        _format_workbook(writer)


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    data_path = project_root / "data" / "processed" / "vendas_processadas.csv"
    output_path = project_root / "output" / "relatorio_vendas.xlsx"

    if not data_path.is_file():
        raise FileNotFoundError(f"Dataset processado não encontrado: {data_path}")
    df = pd.read_csv(data_path)
    if df.empty:
        raise ValueError(f"Dataset processado está vazio: {data_path}")

    kpis = calculate_kpis(df)
    print("Resumo Executivo")
    print(f"Faturamento total: {_format_brl(kpis['faturamento_total'])}")
    print(f"Quantidade total de itens: {_format_int(kpis['quantidade_total'])}")
    print(f"Total de vendas: {_format_int(kpis['total_vendas'])}")
    print(f"Ticket médio: {_format_brl(kpis['ticket_medio'])}")

    generate_excel_report(df, output_path)
    print(f"Relatório Excel gerado: {output_path}")


if __name__ == "__main__":
    main()
